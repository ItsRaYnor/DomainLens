"""
Rapid7 InsightVM / Nexpose vulnerability export import.

Accepts CSV and XLSX exports from:
  - InsightVM / Nexpose CSV Export (Basic Vulnerability Check Results)
  - SQL Query Export / Query Builder CSV
  - Custom Excel exports with similar column names

Column names are matched flexibly (case-insensitive, spaces/underscores ignored).
"""

from __future__ import annotations

import csv
import io
import re
from typing import BinaryIO, Iterable

# Normalized field → accepted header aliases (lowercase, stripped of punctuation)
_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "asset_ip": (
        "ip address", "asset ip address", "asset ip", "ip", "host ip",
        "asset_ip_address", "ip_address", "ipv4 address", "ipv4",
    ),
    "hostname": (
        "host name", "hostname", "asset names", "asset name", "dns name",
        "fqdn", "host_name", "asset_names", "asset_name", "dns_name",
        "name", "asset", "host",
    ),
    "title": (
        "vulnerability title", "title", "vulnerability_title", "vuln title",
        "vulnerability name", "vulnerability_name", "check name", "check_name",
    ),
    "severity": (
        "vulnerability severity level", "severity", "severity level",
        "vulnerability_severity", "severity_level", "risk", "risk score",
        "vulnerability severity", "severity_score",
    ),
    "cvss": (
        "vulnerability cvss score", "cvss", "cvss score", "cvss_score",
        "vulnerability_cvss_score", "cvss v3", "cvssv3", "cvss3",
        "cvss_v3_score", "base score",
    ),
    "cves": (
        "vulnerability cve ids", "cve", "cves", "cve ids", "cve_ids",
        "vulnerability_cves", "vulnerability_cve_ids", "cve id",
    ),
    "solution": (
        "vulnerability solution", "solution", "remediation", "fix",
        "vulnerability_solution", "remediation steps", "solution summary",
    ),
    "description": (
        "vulnerability description", "description", "vulnerability_description",
        "synopsis", "proof", "vulnerability proof",
    ),
    "port": (
        "service port", "port", "port number", "service_port", "port_number",
    ),
    "protocol": (
        "service protocol", "protocol", "service_protocol", "proto",
    ),
    "vuln_id": (
        "vulnerability id", "vuln id", "vulnerability_id", "nexpose id",
        "nexpose_id", "vulnerability check id", "check id", "plugin id",
    ),
    "result_code": (
        "vulnerability result code", "result code", "result_code",
        "vulnerability_result_code",
    ),
    "exploit": (
        "exploit", "exploits", "exploit count", "exploit_count",
        "malware kit", "malware_kits",
    ),
}

_SEVERITY_MAP = {
    "critical": "critical",
    "severe": "high",
    "high": "high",
    "medium": "medium",
    "moderate": "medium",
    "low": "low",
    "info": "info",
    "informational": "info",
    "none": "info",
}


def _norm_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (value or "").strip().lower()).strip()


def _build_header_map(headers: Iterable[str]) -> dict[str, str]:
    """Map normalized field name → original header string present in file."""
    by_norm = {_norm_header(h): h for h in headers if h is not None}
    mapping: dict[str, str] = {}
    for field, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            key = _norm_header(alias)
            if key in by_norm:
                mapping[field] = by_norm[key]
                break
    return mapping


def _cell(row: dict, header_map: dict[str, str], field: str, default: str = "") -> str:
    header = header_map.get(field)
    if not header:
        return default
    val = row.get(header)
    if val is None:
        return default
    return str(val).strip()


def _parse_cvss(raw: str):
    if not raw:
        return None
    try:
        return float(str(raw).replace(",", ".").strip())
    except ValueError:
        return None


def _parse_port(raw: str):
    if not raw:
        return None
    try:
        port = int(float(str(raw).strip()))
        return port if 0 < port <= 65535 else None
    except ValueError:
        return None


def _parse_cves(raw: str) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"[,;\s|/]+", raw.strip())
    out = []
    for p in parts:
        p = p.strip().upper()
        if re.match(r"^CVE-\d{4}-\d{4,}$", p):
            out.append(p)
    return sorted(set(out))


def normalize_severity(raw: str, cvss: float | None = None) -> str:
    text = (raw or "").strip().lower()
    if text in _SEVERITY_MAP:
        return _SEVERITY_MAP[text]
    # Numeric severity / risk score from Nexpose (0–10 style)
    try:
        num = float(text.replace(",", "."))
        return severity_from_cvss(num)
    except ValueError:
        pass
    if cvss is not None:
        return severity_from_cvss(cvss)
    return "info"


def severity_from_cvss(score: float) -> str:
    if score >= 9.0:
        return "critical"
    if score >= 7.0:
        return "high"
    if score >= 4.0:
        return "medium"
    if score > 0:
        return "low"
    return "info"


def extract_domain(hostname: str, asset_ip: str = "") -> str | None:
    """Best-effort registrable domain / FQDN for correlation with DomainLens scans."""
    host = (hostname or "").strip().lower()
    # Asset Names can be comma-separated
    if "," in host:
        host = host.split(",", 1)[0].strip()
    host = host.rstrip(".")
    if not host or re.match(r"^\d{1,3}(\.\d{1,3}){3}$", host):
        return None
    if " " in host:
        host = host.split()[0]
    # Strip wildcards
    host = host.lstrip("*.")
    if not host or "." not in host:
        return None
    return host


def row_to_finding(row: dict, header_map: dict[str, str]) -> dict | None:
    title = _cell(row, header_map, "title")
    hostname = _cell(row, header_map, "hostname")
    asset_ip = _cell(row, header_map, "asset_ip")
    if not title and not hostname and not asset_ip:
        return None
    if not title:
        title = "Untitled vulnerability"

    cvss = _parse_cvss(_cell(row, header_map, "cvss"))
    severity = normalize_severity(_cell(row, header_map, "severity"), cvss)
    cves = _parse_cves(_cell(row, header_map, "cves"))
    domain = extract_domain(hostname, asset_ip)

    return {
        "asset_ip": asset_ip or None,
        "hostname": hostname or None,
        "domain": domain,
        "title": title,
        "severity": severity,
        "cvss": cvss,
        "cves": cves,
        "solution": _cell(row, header_map, "solution") or None,
        "description": _cell(row, header_map, "description") or None,
        "port": _parse_port(_cell(row, header_map, "port")),
        "protocol": _cell(row, header_map, "protocol") or None,
        "vuln_id": _cell(row, header_map, "vuln_id") or None,
        "result_code": _cell(row, header_map, "result_code") or None,
        "exploit": _cell(row, header_map, "exploit") or None,
        "raw": {k: (None if v is None else str(v)) for k, v in row.items()},
    }


def _rows_from_csv(data: bytes | str) -> list[dict]:
    if isinstance(data, bytes):
        text = data.decode("utf-8-sig", errors="replace")
    else:
        text = data
    # Sniff delimiter
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [dict(row) for row in reader]


def _rows_from_xlsx(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "XLSX support requires openpyxl. Install with: pip install openpyxl"
        ) from exc

    wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = next(rows_iter)
    except StopIteration:
        return []
    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(headers)]
    out = []
    for values in rows_iter:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        row = {}
        for i, header in enumerate(headers):
            row[header] = values[i] if i < len(values) else None
        out.append(row)
    return out


def detect_format(filename: str | None, content_type: str | None = None) -> str:
    name = (filename or "").lower()
    ctype = (content_type or "").lower()
    if name.endswith(".xlsx") or "spreadsheetml" in ctype or "excel" in ctype:
        return "xlsx"
    if name.endswith(".xls"):
        raise ValueError("Legacy .xls is not supported; save as .xlsx or .csv")
    return "csv"


def parse_export(
    data: bytes,
    *,
    filename: str | None = None,
    content_type: str | None = None,
    max_rows: int = 50000,
) -> dict:
    """
    Parse a Rapid7 export file into normalized findings.

    Returns:
      {
        source_format, headers, header_map, row_count, findings, summary, warnings
      }
    """
    fmt = detect_format(filename, content_type)
    if fmt == "xlsx":
        rows = _rows_from_xlsx(data)
    else:
        rows = _rows_from_csv(data)

    warnings = []
    if not rows:
        return {
            "source_format": fmt,
            "headers": [],
            "header_map": {},
            "row_count": 0,
            "findings": [],
            "summary": {"by_severity": {}, "assets": 0, "with_cve": 0},
            "warnings": ["No data rows found in export."],
        }

    headers = list(rows[0].keys())
    header_map = _build_header_map(headers)
    if "title" not in header_map and "hostname" not in header_map and "asset_ip" not in header_map:
        warnings.append(
            "Could not recognize Rapid7 columns. Expected at least Vulnerability Title "
            "or Host Name / IP Address. Check that this is an InsightVM/Nexpose export."
        )

    if len(rows) > max_rows:
        warnings.append(f"Truncated import to first {max_rows} rows (file had {len(rows)}).")
        rows = rows[:max_rows]

    findings = []
    for row in rows:
        finding = row_to_finding(row, header_map)
        if finding:
            findings.append(finding)

    by_sev: dict[str, int] = {}
    assets = set()
    with_cve = 0
    for f in findings:
        by_sev[f["severity"]] = by_sev.get(f["severity"], 0) + 1
        key = (f.get("hostname") or "") + "|" + (f.get("asset_ip") or "")
        assets.add(key)
        if f.get("cves"):
            with_cve += 1

    return {
        "source_format": fmt,
        "headers": headers,
        "header_map": header_map,
        "row_count": len(rows),
        "findings": findings,
        "summary": {
            "by_severity": by_sev,
            "assets": len(assets),
            "with_cve": with_cve,
            "findings": len(findings),
        },
        "warnings": warnings,
    }


def findings_for_domain(findings: list[dict], domain: str) -> list[dict]:
    """Filter findings that match a scanned domain (exact host or subdomain)."""
    target = (domain or "").strip().lower().rstrip(".")
    if not target:
        return []
    matched = []
    for f in findings:
        host = (f.get("domain") or f.get("hostname") or "").strip().lower().rstrip(".")
        if not host:
            continue
        if host == target or host.endswith("." + target):
            matched.append(f)
    return matched


def finding_to_advies(finding: dict, domain: str | None = None) -> dict:
    """Convert one Rapid7 finding into an Advies recommendation item."""
    cves = finding.get("cves") or []
    cve_txt = ", ".join(cves[:5]) if cves else ""
    title = finding.get("title") or "Rapid7 vulnerability"
    if cve_txt:
        title = f"{title} ({cve_txt})"

    host = finding.get("hostname") or finding.get("asset_ip") or domain or "asset"
    port = finding.get("port")
    loc = host
    if port:
        loc = f"{host}:{port}"

    problem_parts = [
        f"InsightVM/Nexpose reported this on `{loc}`.",
    ]
    if finding.get("description"):
        problem_parts.append(str(finding["description"])[:400])
    if finding.get("cvss") is not None:
        problem_parts.append(f"CVSS: {finding['cvss']}.")
    if finding.get("exploit"):
        problem_parts.append(f"Exploit/malware signal: {finding['exploit']}.")

    fix = finding.get("solution") or (
        "Apply the vendor patch or configuration change recommended in InsightVM, "
        "then re-scan the asset in Rapid7."
    )
    retest = (
        f"Re-scan `{loc}` in InsightVM/Nexpose and confirm the vulnerability is gone. "
        f"Optionally re-run DomainLens for `{domain or host}` after DNS/TLS/web changes."
    )
    ref = None
    if cves:
        ref = f"https://nvd.nist.gov/vuln/detail/{cves[0]}"

    return {
        "severity": finding.get("severity") or "info",
        "category": "rapid7",
        "title": title[:200],
        "problem": " ".join(problem_parts),
        "fix": fix[:2000],
        "retest": retest,
        "reference": ref,
        "source": "rapid7",
        "vuln_id": finding.get("vuln_id"),
        "cves": cves,
    }


def build_scan_section(findings: list[dict], domain: str) -> dict:
    """Build results['rapid7'] payload for a domain scan."""
    matched = findings_for_domain(findings, domain)
    by_sev: dict[str, int] = {}
    for f in matched:
        by_sev[f["severity"]] = by_sev.get(f["severity"], 0) + 1
    return {
        "success": True,
        "source": "insightvm_nexpose",
        "domain": domain,
        "finding_count": len(matched),
        "by_severity": by_sev,
        "findings": matched[:200],
        "truncated": len(matched) > 200,
    }
